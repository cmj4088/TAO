#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
提取 .docx 文件中所有表格内容和格式信息，输出为结构化 JSON，供 AI 分析。
用法：python extract_docx.py <文件路径>
"""

import sys
import json
import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def get_alignment(paragraph):
    """获取段落对齐方式"""
    align_map = {
        WD_ALIGN_PARAGRAPH.LEFT: "左对齐",
        WD_ALIGN_PARAGRAPH.CENTER: "居中",
        WD_ALIGN_PARAGRAPH.RIGHT: "右对齐",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "两端对齐",
    }
    return align_map.get(paragraph.alignment, "默认")


def get_run_format(run):
    """提取 run 的格式信息"""
    fmt = {}
    font = run.font
    if font.name:
        fmt["字体"] = font.name
    if font.size:
        fmt["字号"] = f"{font.size.pt}pt"
    if font.bold:
        fmt["加粗"] = True
    if font.italic:
        fmt["斜体"] = True
    if font.underline:
        fmt["下划线"] = True
    if font.color and font.color.rgb:
        fmt["颜色"] = str(font.color.rgb)
    return fmt


def extract_cell_info(cell):
    """提取单个单元格的完整信息（包含格式）"""
    paragraphs_info = []
    full_text_parts = []

    for para in cell.paragraphs:
        para_info = {
            "对齐": get_alignment(para),
            "文本片段": [],
        }
        para_text = ""
        for run in para.runs:
            run_info = {"文本": run.text, "格式": get_run_format(run)}
            para_info["文本片段"].append(run_info)
            para_text += run.text
        para_info["完整文本"] = para_text
        paragraphs_info.append(para_info)
        full_text_parts.append(para_text)

    return {
        "完整文本": "\n".join(full_text_parts),
        "段落详情": paragraphs_info,
    }


def extract_tables(doc):
    """提取文档中所有表格"""
    tables = []

    for idx, table in enumerate(doc.tables):
        table_info = {
            "表格编号": idx + 1,
            "行数": len(table.rows),
            "列数": len(table.columns),
            "数据": [],
        }

        # 收集每个单元格的字体信息，用于统计
        all_fonts = {}

        for row_idx, row in enumerate(table.rows):
            row_data = []
            for col_idx, cell in enumerate(row.cells):
                cell_info = extract_cell_info(cell)
                row_data.append(cell_info)

                # 统计字体使用情况
                for para in cell.paragraphs:
                    for run in para.runs:
                        font_name = run.font.name or "默认字体"
                        font_size = run.font.size.pt if run.font.size else "默认字号"
                        key = f"{font_name} / {font_size}pt"
                        if key not in all_fonts:
                            all_fonts[key] = 0
                        all_fonts[key] += 1

            table_info["数据"].append({"行号": row_idx + 1, "单元格": row_data})

        table_info["字体统计"] = all_fonts
        table_info["使用的字体种类数"] = len(
            set(k.split(" / ")[0] for k in all_fonts.keys())
        )

        tables.append(table_info)

    return tables


def main():
    if len(sys.argv) < 2:
        print("用法：python extract_docx.py <docx文件路径>")
        sys.exit(1)

    filepath = sys.argv[1]
    if not os.path.exists(filepath):
        print(f"文件不存在：{filepath}")
        sys.exit(1)

    doc = Document(filepath)

    result = {
        "文件名": os.path.basename(filepath),
        "表格数量": len(doc.tables),
        "表格详情": extract_tables(doc),
    }

    # 输出到同目录下的 JSON 文件
    out_path = os.path.splitext(filepath)[0] + "_分析结果.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"提取完成！共 {len(doc.tables)} 个表格")
    print(f"结果已保存到：{out_path}")
    print()
    # 同时打印简要摘要
    for t in result["表格详情"]:
        print(f"--- 表格 {t['表格编号']}：{t['行数']}行 × {t['列数']}列 ---")
        print(f"    字体种类：{t['使用的字体种类数']} 种")
        for row in t["数据"][:3]:  # 只显示前3行预览
            texts = [c["完整文本"][:20] for c in row["单元格"]]
            print(f"    行{row['行号']}：{' | '.join(texts)}")
        if t["行数"] > 3:
            print(f"    ... 还有 {t['行数'] - 3} 行")
        print()


if __name__ == "__main__":
    main()
