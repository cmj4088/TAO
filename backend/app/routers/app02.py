"""app02 文件审查 API"""
import io
import os
import re
import uuid
import json
import asyncio
import threading
import tempfile
import traceback
from collections import defaultdict
from urllib.parse import quote

import openpyxl
from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse

from app.schemas.app02 import (
    ReviewIssue,
    FileReviewResult,
    UploadedFileInfo,
    UploadResponse,
    ReviewResponse,
)
from app.services.doc_reviewer import classify_template
from app.services.doc_converter import convert_doc_to_docx, convert_docx_to_pdf
from app.services.ai_reviewer_app02 import review_stream
from app.config import ConfigManager

router = APIRouter(prefix="/api/app02", tags=["app02"])

TEMP_DIR = os.path.join(tempfile.gettempdir(), "tao_app02")
os.makedirs(TEMP_DIR, exist_ok=True)

_session_files: dict[str, dict] = {}
_session_results: list[FileReviewResult] = []
_reviewed_file_ids: set[str] = set()

_ai_tasks: dict[str, dict] = {}
_stream_queues: dict[str, dict[str, asyncio.Queue]] = defaultdict(dict)


@router.get("/status")
def get_status():
    return {"status": "ok", "message": "app02 文件审查模块就绪"}


@router.post("/upload", response_model=UploadResponse)
async def upload_files(files: list[UploadFile] = File(...)):
    global _session_files

    infos: list[UploadedFileInfo] = []

    for f in files:
        if not f.filename:
            continue
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in (".doc", ".docx"):
            continue

        file_id = str(uuid.uuid4())
        safe_name = f"{file_id}{ext}"
        filepath = os.path.join(TEMP_DIR, safe_name)

        content = await f.read()
        with open(filepath, "wb") as fh:
            fh.write(content)

        is_doc = ext == ".doc"
        file_type = classify_template(f.filename)

        _session_files[file_id] = {
            "filename": f.filename,
            "filepath": filepath,
            "file_type": file_type,
            "is_doc": is_doc,
            "converted_path": None,
        }

        infos.append(UploadedFileInfo(
            file_id=file_id,
            filename=f.filename,
            file_type=file_type,
            is_doc=is_doc,
            converted=False,
        ))

    return UploadResponse(files=infos)


@router.post("/reset")
def reset_session():
    global _session_files, _session_results, _reviewed_file_ids, _ai_tasks, _stream_queues
    _session_files.clear()
    _session_results.clear()
    _reviewed_file_ids.clear()
    _ai_tasks.clear()
    _stream_queues.clear()
    return {"status": "ok"}


@router.post("/review", response_model=ReviewResponse)
def run_review():
    global _session_files, _session_results, _reviewed_file_ids

    if not _session_files:
        raise HTTPException(400, "请先上传文件")

    new_file_ids = [fid for fid in _session_files if fid not in _reviewed_file_ids]
    if not new_file_ids:
        raise HTTPException(400, "所有文件已审查完毕，没有新增文件需要审查")

    new_results: list[FileReviewResult] = []

    for file_id in new_file_ids:
        info = _session_files[file_id]
        filepath = info["filepath"]

        if info["is_doc"]:
            converted = convert_doc_to_docx(filepath)
            if converted:
                info["converted_path"] = converted
                filepath = converted
            else:
                result = FileReviewResult(
                    file_id=file_id,
                    filename=info["filename"],
                    file_type=info["file_type"],
                    teacher="未知",
                    passed=False,
                    fonts_used=[],
                    issues=[ReviewIssue(
                        type="转换失败",
                        detail="无法将 .doc 文件转换为 .docx，请确保安装了 LibreOffice 或 Microsoft Word",
                        location="",
                        snippet="",
                        severity="error",
                    )],
                )
                new_results.append(result)
                _reviewed_file_ids.add(file_id)
                continue

        result = FileReviewResult(
            file_id=file_id,
            filename=info["filename"],
            file_type=info["file_type"],
            teacher="未知",
            passed=True,
            fonts_used=[],
            issues=[],
        )
        new_results.append(result)
        _reviewed_file_ids.add(file_id)

    _session_results = new_results

    ai_task_id = str(uuid.uuid4())
    _ai_tasks[ai_task_id] = {"status": "starting", "total": len(new_file_ids), "completed": 0, "files": [], "findings": [], "error": None}

    mode = "single" if len(_session_results) == 1 else "batch"
    summary = f"共 {len(_session_results)} 个文件，AI 审查已启动，请等待分析完成"

    t = threading.Thread(
        target=_run_ai_review_task,
        args=(ai_task_id, new_file_ids),
        daemon=True,
    )
    t.start()

    return ReviewResponse(mode=mode, results=_session_results, summary=summary, ai_task_id=ai_task_id)


@router.get("/preview/{file_id}")
def preview_file(file_id: str):
    global _session_files

    if file_id not in _session_files:
        raise HTTPException(404, "文件不存在")

    info = _session_files[file_id]
    src_path = info.get("converted_path") or info["filepath"]

    pdf_path = convert_docx_to_pdf(src_path)
    if not pdf_path:
        raise HTTPException(500, "无法转换为 PDF，请确保安装了 LibreOffice 或 Microsoft Word")

    def iterfile():
        with open(pdf_path, "rb") as f:
            yield from f
        try:
            os.remove(pdf_path)
        except Exception:
            pass

    encoded_filename = quote(f"{info['filename']}.pdf")
    return StreamingResponse(
        iterfile(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"inline; filename*=UTF-8''{encoded_filename}"},
    )


@router.get("/export")
def export_excel():
    global _session_results

    if not _session_results:
        raise HTTPException(400, "请先执行审查")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "审查结果汇总"

    headers = ["文件名", "老师", "文件类型", "是否通过", "错误类型", "位置", "问题描述"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)

    row_idx = 2
    for result in _session_results:
        if result.issues:
            for issue in result.issues:
                ws.cell(row=row_idx, column=1, value=result.filename)
                ws.cell(row=row_idx, column=2, value=result.teacher)
                ws.cell(row=row_idx, column=3, value=result.file_type)
                ws.cell(row=row_idx, column=4, value="通过" if result.passed else "不通过")
                ws.cell(row=row_idx, column=5, value=issue.type)
                ws.cell(row=row_idx, column=6, value=issue.location)
                ws.cell(row=row_idx, column=7, value=issue.detail)
                row_idx += 1
        else:
            ws.cell(row=row_idx, column=1, value=result.filename)
            ws.cell(row=row_idx, column=2, value=result.teacher)
            ws.cell(row=row_idx, column=3, value=result.file_type)
            ws.cell(row=row_idx, column=4, value="通过")
            ws.cell(row=row_idx, column=5, value="无")
            row_idx += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    encoded_filename = quote("文件审查结果.xlsx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


# ===== AI 审查（异步并发流式） =====

def _run_ai_review_task(task_id: str, file_ids: list[str]):
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    try:
        loop.run_until_complete(_run_ai_review_task_async(task_id, file_ids))
    except Exception:
        traceback.print_exc()
    finally:
        loop.close()


async def _run_ai_review_task_async(task_id: str, file_ids: list[str]):
    global _session_files, _ai_tasks

    concurrency = int(ConfigManager.get("app02_ai_concurrency", "1"))
    semaphore = asyncio.Semaphore(max(1, concurrency))

    all_files = []
    for fid in file_ids:
        info = _session_files.get(fid)
        if info:
            all_files.append({
                "file_id": fid,
                "filename": info["filename"],
                "status": "waiting",
                "findings": [],
            })

    all_findings: list[dict] = []
    for fe in all_files:
        for f in fe.get("findings", []):
            f.setdefault("file_id", fe["file_id"])
            f.setdefault("filename", fe["filename"])
        all_findings.extend(fe.get("findings", []))

    _ai_tasks[task_id] = {
        "status": "running",
        "total": len(all_files),
        "completed": sum(1 for fe in all_files if fe["status"] == "done"),
        "current_file_id": None,
        "files": all_files,
        "findings": all_findings,
        "error": None,
    }

    file_queues: dict[str, asyncio.Queue] = {fid: asyncio.Queue() for fid in file_ids}
    _stream_queues[task_id] = file_queues

    async def review_one_file(fe: dict):
        fid = fe["file_id"]
        info = _session_files.get(fid)
        if not info:
            fe["status"] = "done"
            _ai_tasks[task_id]["completed"] += 1
            await file_queues[fid].put(("file_done", {"file_id": fid, "findings": [], "extracted_info": {}}))
            return

        filepath = info.get("converted_path") or info["filepath"]
        if not os.path.exists(filepath):
            fe["status"] = "done"
            _ai_tasks[task_id]["completed"] += 1
            await file_queues[fid].put(("file_done", {"file_id": fid, "findings": [], "extracted_info": {}}))
            return

        fe["status"] = "reviewing"
        _ai_tasks[task_id]["current_file_id"] = fid

        try:
            async with semaphore:
                async for event_type, data in review_stream(filepath, info["filename"]):
                    if event_type == "reasoning":
                        await file_queues[fid].put(("reasoning", {"file_id": fid, "content": data}))
                    elif event_type == "token":
                        await file_queues[fid].put(("token", {"file_id": fid, "content": data}))
                    elif event_type == "done":
                        findings, extracted_info, _raw_text = data
                        for f in findings:
                            f["file_id"] = fid
                            f["filename"] = info["filename"]
                        fe["findings"] = findings
                        fe["extracted_info"] = extracted_info
                        all_findings.extend(findings)
                        await file_queues[fid].put(("file_done", {
                            "file_id": fid,
                            "findings": findings,
                            "extracted_info": extracted_info,
                        }))
                    elif event_type == "error":
                        await file_queues[fid].put(("file_error", {"file_id": fid, "error": str(data)}))
        except Exception as e:
            traceback.print_exc()
            fe["findings"] = []
            fe["extracted_info"] = {}
            try:
                await file_queues[fid].put(("file_error", {"file_id": fid, "error": str(e)}))
            except Exception:
                pass

        fe["status"] = "done"
        _ai_tasks[task_id]["completed"] += 1

    tasks_to_review = [fe for fe in all_files if fe["status"] != "done"]
    await asyncio.gather(*(review_one_file(fe) for fe in tasks_to_review))

    _ai_tasks[task_id]["status"] = "done"
    _ai_tasks[task_id]["current_file_id"] = None
    _ai_tasks[task_id]["findings"] = all_findings


@router.get("/ai-review/stream/{task_id}")
async def ai_review_stream_endpoint(task_id: str):
    if task_id not in _stream_queues and task_id not in _ai_tasks:
        raise HTTPException(404, "任务不存在或已过期")

    async def event_generator():
        queues: dict = {}
        file_ids: list = []

        for _ in range(60):
            queues = _stream_queues.get(task_id, {})
            file_ids = list(queues.keys())
            if file_ids:
                break
            await asyncio.sleep(0.5)

        if not file_ids:
            yield f"event: error\ndata: {json.dumps({'error': 'AI 审查启动超时，请重试'})}\n\n"
            yield f"event: task_done\ndata: {json.dumps({'status': 'done'})}\n\n"
            return

        total_files = len(file_ids)

        merged_queue: asyncio.Queue = asyncio.Queue()

        async def forward_from_file(fid: str):
            q = queues.get(fid)
            if not q:
                return
            while True:
                try:
                    event_type, data = await asyncio.wait_for(q.get(), timeout=600.0)
                    await merged_queue.put((event_type, data))
                    if event_type in ("file_done", "file_error"):
                        return
                except asyncio.TimeoutError:
                    await merged_queue.put(("file_error", {"file_id": fid, "error": "审查超时"}))
                    return
                except Exception:
                    return

        forwarders = [asyncio.create_task(forward_from_file(fid)) for fid in file_ids]

        finished = 0
        while finished < total_files:
            try:
                event_type, data = await asyncio.wait_for(merged_queue.get(), timeout=600.0)
            except asyncio.TimeoutError:
                yield f"event: error\ndata: {json.dumps({'error': '整体审查超时'})}\n\n"
                break

            if event_type == "reasoning":
                yield f"event: reasoning\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"
            elif event_type == "token":
                yield f"event: token\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"
            elif event_type == "file_done":
                yield f"event: file_done\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"
                finished += 1
            elif event_type == "file_error":
                yield f"event: file_error\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"
                finished += 1

        yield f"event: task_done\ndata: {json.dumps({'status': 'done'})}\n\n"

        for t in forwarders:
            t.cancel()
        await asyncio.gather(*forwarders, return_exceptions=True)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.get("/ai-review/status")
def ai_review_status():
    api_key = ConfigManager.get("llm_key", "")
    return {"configured": bool(api_key)}


@router.post("/ai-review/start")
def start_ai_review():
    global _session_files
    if not _session_files:
        raise HTTPException(400, "请先上传文件并执行审查")

    task_id = str(uuid.uuid4())
    all_file_ids = list(_session_files.keys())
    t = threading.Thread(target=_run_ai_review_task, args=(task_id, all_file_ids), daemon=True)
    t.start()

    return {"task_id": task_id, "status": "running"}


@router.get("/ai-review/progress/{task_id}")
def ai_review_progress(task_id: str):
    task = _ai_tasks.get(task_id)
    if not task:
        raise HTTPException(404, "任务不存在或已过期")
    return task
