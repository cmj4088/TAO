"""app01 监考分配 API"""
import io
import copy
import json
import uuid
import asyncio
import threading
from collections import defaultdict
from urllib.parse import quote

import openpyxl
from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse

from app.schemas.app01 import (
    TeacherInfo,
    ExamRow,
    UploadResponse,
    AllocateRequest,
    AllocateResponse,
    SwapRequest,
    SwapResponse,
    ReplaceRequest,
    ReplaceResponse,
    SetRowsRequest,
    ValidateResponse,
    ValidationError,
)
from app.services.invigilator import allocate, validate
from app.services.ai_reviewer_app01 import review_stream

router = APIRouter(prefix="/api/app01", tags=["app01"])

# 当前会话状态（内存中保存已上传解析的数据和分配结果）
_session_exam_rows: list[ExamRow] | None = None
_session_teachers: list[TeacherInfo] | None = None
_original_exam_bytes: bytes | None = None  # 原始考试安排文件字节，导出时重新加载


def _extract_name(teacher_str: str) -> str:
    """从 '2022000012/李莉' 提取 '李莉'"""
    if "/" in teacher_str:
        return teacher_str.split("/", 1)[1].strip()
    return teacher_str.strip()


def _parse_exam_schedule(file_content: bytes) -> list[ExamRow]:
    """解析考试安排表"""
    wb = openpyxl.load_workbook(io.BytesIO(file_content))
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:  # 跳过标题行
            continue
        if row[9] is None:  # 无考试时间，跳过空行
            continue

        raw_teachers = str(row[8]) if row[8] else ""
        teacher_names = [
            _extract_name(t.strip())
            for t in raw_teachers.split(",")
            if t.strip()
        ]

        rows.append(ExamRow(
            index=i,
            场次=str(row[1]) if row[1] else "",
            班级名称=str(row[6]) if row[6] else "",
            课程名称=str(row[7]) if row[7] else "",
            任课教师=teacher_names,
            考试时间=str(row[9]).strip() if row[9] else "",
            考试地点=str(row[10]).strip() if row[10] else "",
            人数=str(row[11]) if row[11] else "",
        ))
    return rows


def _parse_contact_list(file_content: bytes) -> list[TeacherInfo]:
    """解析通讯录（新版4列：岗位、姓名、安排场次、分组）"""
    wb = openpyxl.load_workbook(io.BytesIO(file_content))
    ws = wb[wb.sheetnames[0]]
    teachers = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:  # 跳过标题行
            continue

        name = str(row[1]).replace(" ", "").strip() if row[1] else ""
        if not name:
            continue

        department = str(row[0]).strip() if row[0] else ""

        # C列：安排场次
        slots = 0
        if row[2] is not None:
            try:
                slots = int(row[2])
            except (ValueError, TypeError):
                slots = 0

        # D列：分组（保留大小写）
        group = str(row[3]).strip() if row[3] else ""

        teachers.append(TeacherInfo(
            name=name,
            department=department,
            slots=slots,
            group=group,
        ))
    return teachers


@router.get("/status")
def get_status():
    return {"status": "ok", "message": "app01 监考分配模块就绪"}


@router.post("/upload", response_model=UploadResponse)
async def upload_files(
    exam_file: UploadFile = File(...),
    contact_file: UploadFile = File(...),
):
    """上传考试安排表和通讯录，返回解析后的预览数据"""
    global _session_exam_rows, _session_teachers, _original_exam_bytes

    if not exam_file.filename or not exam_file.filename.endswith(".xlsx"):
        raise HTTPException(400, "考试安排表必须是 .xlsx 文件")
    if not contact_file.filename or not contact_file.filename.endswith(".xlsx"):
        raise HTTPException(400, "通讯录必须是 .xlsx 文件")

    exam_content = await exam_file.read()
    contact_content = await contact_file.read()

    _original_exam_bytes = exam_content

    _session_exam_rows = _parse_exam_schedule(exam_content)
    _session_teachers = _parse_contact_list(contact_content)

    return UploadResponse(
        exam_rows=_session_exam_rows,
        teachers=_session_teachers,
    )


@router.post("/allocate", response_model=AllocateResponse)
def run_allocation(body: AllocateRequest):
    """执行监考分配算法"""
    global _session_exam_rows, _session_teachers

    if not body.exam_rows:
        raise HTTPException(400, "考试安排数据为空")

    rows = copy.deepcopy(body.exam_rows)
    teachers = body.teachers if body.teachers else (_session_teachers or [])

    result_rows, warnings, teacher_loads = allocate(rows, teachers)
    _session_exam_rows = result_rows
    _session_teachers = teachers

    return AllocateResponse(exam_rows=result_rows, warnings=warnings, teacher_loads=teacher_loads)


@router.post("/swap", response_model=SwapResponse)
def swap_teacher(body: SwapRequest):
    """交换两个监考格子的老师"""
    global _session_exam_rows

    if not _session_exam_rows:
        raise HTTPException(400, "请先上传文件并执行分配")

    rows = copy.deepcopy(_session_exam_rows)
    src_val = None
    tgt_val = None
    for row in rows:
        if row.index == body.source_row_index:
            src_val = getattr(row, body.source_position)
        if row.index == body.target_row_index:
            tgt_val = getattr(row, body.target_position)

    for row in rows:
        if row.index == body.source_row_index:
            setattr(row, body.source_position, tgt_val)
        if row.index == body.target_row_index:
            setattr(row, body.target_position, src_val)

    _session_exam_rows = rows
    return SwapResponse(exam_rows=rows)


@router.post("/replace", response_model=ReplaceResponse)
def replace_teacher(body: ReplaceRequest):
    """将指定格子的监考老师替换为新老师"""
    global _session_exam_rows

    if not _session_exam_rows:
        raise HTTPException(400, "请先上传文件并执行分配")

    rows = copy.deepcopy(_session_exam_rows)
    found = False
    for row in rows:
        if row.index == body.row_index:
            setattr(row, body.position, body.new_teacher if body.new_teacher else None)
            found = True
            break

    if not found:
        raise HTTPException(400, f"未找到行 index={body.row_index}")

    _session_exam_rows = rows
    return ReplaceResponse(exam_rows=rows)


@router.post("/set-rows", response_model=list[ExamRow])
def set_rows(body: SetRowsRequest):
    """批量替换所有考试行（用于撤销等场景）"""
    global _session_exam_rows
    _session_exam_rows = copy.deepcopy(body.exam_rows)
    return _session_exam_rows


@router.post("/validate", response_model=ValidateResponse)
def run_validation(body: AllocateRequest | None = None):
    """校验当前分配结果，返回违规项列表"""
    rows = _session_exam_rows
    if not rows:
        raise HTTPException(400, "请先上传文件并执行分配")

    teachers = _session_teachers or []
    if body and body.teachers:
        teachers = body.teachers
    if body and body.exam_rows:
        rows = body.exam_rows

    error_dicts = validate(rows, teachers)
    errors = [ValidationError(**e) for e in error_dicts]
    return ValidateResponse(errors=errors)


@router.get("/export")
def export_excel():
    """导出填充后的 Excel"""
    global _session_exam_rows, _original_exam_bytes

    if not _session_exam_rows or not _original_exam_bytes:
        raise HTTPException(400, "请先上传文件并执行分配")

    wb = openpyxl.load_workbook(io.BytesIO(_original_exam_bytes))
    ws = wb[wb.sheetnames[0]]

    index_map: dict[int, ExamRow] = {r.index: r for r in _session_exam_rows}

    for row_idx in range(3, ws.max_row + 1):
        r = index_map.get(row_idx - 1)
        if r is None:
            continue
        ws.cell(row=row_idx, column=13).value = r.监考1
        ws.cell(row=row_idx, column=15).value = r.监考2

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    encoded_filename = quote("监考安排结果.xlsx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
        },
    )


# ===== AI 审查（SSE 流式） =====

_ai_tasks: dict[str, dict] = {}
_stream_queues: dict[str, dict[str, asyncio.Queue]] = defaultdict(dict)

TASK_TTL = 600  # 任务完成后保留 10 分钟再清理


def _cleanup_task(task_id: str):
    """延迟清理任务数据，避免内存泄漏"""
    import time
    time.sleep(TASK_TTL)
    _ai_tasks.pop(task_id, None)
    _stream_queues.pop(task_id, None)


@router.post("/ai-review")
def start_ai_review():
    """启动 AI 审查，返回 task_id"""
    global _session_exam_rows, _session_teachers

    if not _session_exam_rows:
        raise HTTPException(400, "请先上传文件并执行分配")

    task_id = str(uuid.uuid4())
    _ai_tasks[task_id] = {"status": "starting"}

    loads: dict[str, int] = {}
    for r in _session_exam_rows:
        if r.监考1:
            loads[r.监考1] = loads.get(r.监考1, 0) + 1
        if r.监考2:
            loads[r.监考2] = loads.get(r.监考2, 0) + 1

    t = threading.Thread(
        target=_run_ai_review,
        args=(task_id, copy.deepcopy(_session_exam_rows), copy.deepcopy(_session_teachers or []), loads),
        daemon=True,
    )
    t.start()

    return {"task_id": task_id, "status": "running"}


def _run_ai_review(task_id: str, exam_rows: list, teachers: list, loads: dict):
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    try:
        loop.run_until_complete(_run_ai_review_async(task_id, exam_rows, teachers, loads))
    except Exception:
        import traceback
        traceback.print_exc()
    finally:
        loop.close()
        threading.Thread(target=_cleanup_task, args=(task_id,), daemon=True).start()


async def _run_ai_review_async(task_id: str, exam_rows: list, teachers: list, loads: dict):
    queue: asyncio.Queue = asyncio.Queue()
    _stream_queues[task_id] = {"app01": queue}

    _ai_tasks[task_id] = {"status": "running"}

    try:
        async for event_type, data in review_stream(exam_rows, teachers, loads):
            await queue.put((event_type, data))
    except Exception as e:
        await queue.put(("error", str(e)))

    _ai_tasks[task_id]["status"] = "done"


@router.get("/ai-review/stream/{task_id}")
async def ai_review_stream_endpoint(task_id: str):
    if task_id not in _stream_queues and task_id not in _ai_tasks:
        raise HTTPException(404, "任务不存在或已过期")

    async def event_generator():
        for _ in range(60):
            queues = _stream_queues.get(task_id, {})
            if queues:
                break
            await asyncio.sleep(0.5)

        q = queues.get("app01")
        if not q:
            yield f"event: review_error\ndata: {json.dumps({'error': 'AI审查启动超时，请重试'})}\n\n"
            yield f"event: task_done\ndata: {json.dumps({'status': 'done'})}\n\n"
            return

        while True:
            try:
                event_type, data = await asyncio.wait_for(q.get(), timeout=300.0)
            except asyncio.TimeoutError:
                yield f"event: review_error\ndata: {json.dumps({'error': '审查超时，请重试'})}\n\n"
                break

            if event_type == "reasoning":
                yield f"event: reasoning\ndata: {json.dumps({'content': data}, ensure_ascii=False)}\n\n"
            elif event_type == "token":
                yield f"event: token\ndata: {json.dumps({'content': data}, ensure_ascii=False)}\n\n"
            elif event_type == "done":
                findings, summary = data
                yield f"event: done\ndata: {json.dumps({'findings': findings, 'summary': summary}, ensure_ascii=False)}\n\n"
                break
            elif event_type == "error":
                yield f"event: review_error\ndata: {json.dumps({'error': str(data)})}\n\n"
                break

        yield f"event: task_done\ndata: {json.dumps({'status': 'done'})}\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )
