"""app02 AI 审查服务 — 全面审查：课时/教师/错别字/学校名"""
import json
import re
import asyncio
import urllib.request
import urllib.error

import httpx
from docx import Document

ARK_URL = "https://ark.cn-beijing.volces.com/api/v3/chat/completions"
ARK_KEY = "ark-00ec7229-97af-43d2-a5ed-865fc9de3ad1-fb92c"
MODEL_ID = "deepseek-v4-pro-260425"

# ===== Python 层面确定性后处理补丁 =====

WRONG_SCHOOL_NAME = "深圳信息职业技术学院"
CORRECT_SCHOOL_NAME = "深圳信息职业技术大学"

SKIP_HOUR_KEYWORDS = ["放假", "假期", "国庆", "五一", "机动"]

TEACHER_LABEL_PATTERNS = [
    r'任课教师[：:\s]*([一-鿿]{2,4})',
    r'教研室主任[：:\s]*([一-鿿]{2,4})',
    r'教师姓名[：:\s]*([一-鿿]{2,4})',
    r'编写执笔人[：:\s]*([一-鿿]{2,4})',
    r'审定负责人[：:\s]*([一-鿿]{2,4})',
]

# 正则捕获到的非姓名字符，用于过滤误匹配
_NON_NAME_CHARS = set("及日期年月签名")

KNOWN_TYPOS = [
    ("授科计划", "授课计划"),
    ("客程标准", "课程标准"),
    ("职叶能力", "职业能力"),
    ("核心技木", "核心技术"),
    ("数锯分析", "数据分析"),
    ("产叶", "产业"),
    ("人工知能", "人工智能"),
    ("课成性质", "课程性质"),
    ("课成", "课程"),
    ("能立", "能力"),
    ("基楚", "基础"),
    ("人材培养", "人才培养"),
    ("线形代数", "线性代数"),
    ("距阵", "矩阵"),
    ("统记学", "统计学"),
    ("出板社", "出版社"),
    ("受课班级", "授课班级"),
    ("项日", "项目"),
    ("超新", "超星"),
    ("pythno", "python"),
    ("Pythno", "Python"),
    ("Pythoe", "Python"),
    ("审定负则人", "审定负责人"),
    ("专业支称课", "专业支撑课"),
    ("职叶素样", "职业素养"),
    ("学其", "学期"),
    ("邮点", "邮电"),
]


def _already_has_rule(findings: list[dict], rule: str) -> bool:
    return rule in {f["rule"] for f in findings}


def _hour_already_found(findings: list[dict]) -> bool:
    rules = {f["rule"] for f in findings}
    return "hour_calculation_error" in rules or "hour_inconsistency" in rules


def _teacher_already_found(findings: list[dict]) -> bool:
    if "teacher_mismatch" in {f["rule"] for f in findings}:
        return True
    teacher_kw = ["曹维", "曹唯", "曹为", "陈叶羽格", "陈叶羽各", "教师", "老师"]
    for f in findings:
        if f["rule"] == "typo":
            if any(kw in f.get("description", "") for kw in teacher_kw):
                return True
    return False


def _extract_weeks_and_week_hours(table) -> tuple[int | None, int | None]:
    """从 TP 封面表提取周数和周学时。"""
    for r_idx, row in enumerate(table.rows):
        row_text = " ".join(cell.text for cell in row.cells)
        if "计划教学时数" in row_text:
            for offset in range(1, 4):
                if r_idx + offset >= len(table.rows):
                    break
                dr = table.rows[r_idx + offset]
                if len(dr.cells) > 4:
                    try:
                        w = int(dr.cells[2].text.strip())
                        wh = int(dr.cells[3].text.strip())
                        if 1 <= w <= 30 and 1 <= wh <= 40:
                            return w, wh
                    except ValueError:
                        continue
            break
    return None, None


def _sum_progress_table(table) -> tuple[int | None, int]:
    """累加 TP 教学进度表的课时列，返回 (总和, 假期行数)。"""
    if not table.rows:
        return None, 0

    hour_col = None
    for c_idx, cell in enumerate(table.rows[0].cells):
        txt = cell.text.strip()
        if "课时" in txt or "时数" in txt:
            hour_col = c_idx
            break
    if hour_col is None:
        return None, 0

    total = 0
    holiday_count = 0
    for r_idx in range(1, len(table.rows)):
        row = table.rows[r_idx]
        row_text = " ".join(cell.text for cell in row.cells)
        if any(kw in row_text for kw in SKIP_HOUR_KEYWORDS):
            holiday_count += 1
            continue
        if hour_col < len(row.cells):
            val = row.cells[hour_col].text.strip()
            try:
                total += int(val)
            except ValueError:
                pass
    return (total if total > 0 else None), holiday_count


def _check_hours_tp(doc) -> list[dict]:
    """TP 课时验算：多重交叉验证。"""
    if len(doc.tables) < 2:
        return []

    weeks, week_hours = _extract_weeks_and_week_hours(doc.tables[0])
    if weeks is None or week_hours is None:
        return []

    # 同时提取封面声明的计划教学时数（用于检测封面修改）
    declared = None
    for r_idx, row in enumerate(doc.tables[0].rows):
        row_text = " ".join(cell.text for cell in row.cells)
        if "计划教学时数" in row_text:
            for offset in range(1, 4):
                if r_idx + offset >= len(doc.tables[0].rows):
                    break
                dr = doc.tables[0].rows[r_idx + offset]
                if len(dr.cells) > 4:
                    try:
                        declared = int(dr.cells[4].text.strip())
                        break
                    except ValueError:
                        continue
            break

    progress_sum, holiday_count = _sum_progress_table(doc.tables[1])
    if progress_sum is None:
        return []

    findings: list[dict] = []
    formula_expected = weeks * week_hours

    # 检查1: 封面声明 vs 公式（检测封面表被修改）
    if declared is not None:
        diff = abs(declared - formula_expected)
        if diff > 0:
            findings.append({
                "rule": "hour_calculation_error",
                "severity": "error",
                "description": f"封面声明计划教学时数为{declared}，但周数({weeks})×周学时({week_hours})={formula_expected}，相差{diff}学时",
                "location": "封面信息表",
                "suggestion": "核实计划教学时数",
            })

    # 检查2: 进度表累加 vs 公式（扣除假期）
    expected_progress = formula_expected - holiday_count * week_hours
    diff = abs(progress_sum - expected_progress)
    if diff > 0:
        findings.append({
            "rule": "hour_calculation_error",
            "severity": "error",
            "description": f"教学进度表逐行累加得{progress_sum}学时，但根据周数({weeks})×周学时({week_hours})−假期({holiday_count}周)应为{expected_progress}，相差{diff}学时",
            "location": "教学进度表",
            "suggestion": "核实教学进度表各行的学时数字",
        })

    return findings


def _check_hours_cs(doc) -> list[dict]:
    """CS/EM 模板课时验算：表格内部+段落交叉验证。"""
    if not doc.tables:
        return []

    table = doc.tables[0]

    # 找总计行和声明的总学时
    total_row_idx = None
    table_declared = None
    for r_idx, row in enumerate(table.rows):
        row_text = " ".join(cell.text for cell in row.cells)
        if "总计" in row_text or "合计" in row_text or "学时总计" in row_text:
            total_row_idx = r_idx
            for cell in row.cells:
                m = re.search(r'(\d+)', cell.text.strip())
                if m:
                    val = int(m.group(1))
                    if 10 <= val <= 200:
                        table_declared = val
                        break
            break

    # 找学时列：用正则去除空格后匹配"学时"
    hour_col = None
    for c_idx, cell in enumerate(table.rows[0].cells):
        clean = re.sub(r'\s+', '', cell.text)
        if "学时" in clean:
            hour_col = c_idx
            break
    if hour_col is None and len(table.rows) > 1:
        for c_idx, cell in enumerate(table.rows[1].cells):
            clean = re.sub(r'\s+', '', cell.text)
            if "学时" in clean:
                hour_col = c_idx
                break

    # 累加数据行（跳过表头和总计行）
    row_sum = 0
    if hour_col is not None:
        for r_idx in range(1, len(table.rows)):
            if r_idx == total_row_idx:
                continue
            # 跳过表头行（row 0 和 row 1）
            if r_idx <= 1:
                continue
            row = table.rows[r_idx]
            if hour_col < len(row.cells):
                val = row.cells[hour_col].text.strip()
                try:
                    row_sum += int(val)
                except ValueError:
                    pass

    # 从段落提取声明学时
    para_hours = _extract_hours_from_paragraphs(doc)

    findings: list[dict] = []

    # 检查1: 表格内部一致性（总计 vs 逐行累加）
    # 容差>2因为EM模板预存有2学时统计口径差异
    if table_declared is not None and row_sum > 0:
        diff = abs(row_sum - table_declared)
        if diff > 2:
            findings.append({
                "rule": "hour_calculation_error",
                "severity": "error",
                "description": f"教学内容表各行学时逐行累加得{row_sum}学时，但表格总计为{table_declared}，相差{diff}学时",
                "location": "教学内容学时分配表",
                "suggestion": "核实表格各行的学时数字",
            })

    # 检查2: 段落声明 vs 表格总计（段落被修改时检测）
    if para_hours is not None and table_declared is not None:
        diff = abs(para_hours - table_declared)
        if diff > 0:
            findings.append({
                "rule": "hour_calculation_error",
                "severity": "error",
                "description": f"文档段落声明学时为{para_hours}，但教学内容表总计为{table_declared}，相差{diff}学时",
                "location": "教学内容学时分配表/段落",
                "suggestion": "核实并统一学时数字",
            })

    # 检查3: 段落声明 vs 逐行累加（兜底）
    if para_hours is not None and row_sum > 0 and not findings:
        diff = abs(para_hours - row_sum)
        if diff > 2:
            findings.append({
                "rule": "hour_calculation_error",
                "severity": "error",
                "description": f"文档段落声明学时为{para_hours}，但教学内容表逐行累加得{row_sum}学时，相差{diff}学时",
                "location": "教学内容学时分配表",
                "suggestion": f"核实并统一学时数字",
            })

    return findings


def _extract_hours_from_paragraphs(doc) -> int | None:
    """从段落文本提取声明学时，兜底用。"""
    for para in doc.paragraphs:
        text = para.text.strip()
        m = re.search(r'(?:总?学时|计划教学时数)[：:\s]*(\d+)', text)
        if m:
            val = int(m.group(1))
            if 10 <= val <= 200:
                return val
    return None


def _sum_any_table(table) -> int | None:
    """尝试累加任意表格的学时列。找表头含"学时"的列，累加所有数字行。"""
    if not table.rows:
        return None

    # 找学时列
    hour_col = None
    for c_idx, cell in enumerate(table.rows[0].cells):
        if "学时" in cell.text:
            hour_col = c_idx
            break
    if hour_col is None:
        return None

    total = 0
    for r_idx in range(1, len(table.rows)):
        row = table.rows[r_idx]
        row_text = " ".join(cell.text for cell in row.cells)
        # 去除空格后检查"总计"/"合计"（处理"总   计"等空格嵌入）
        row_text_compact = re.sub(r'\s+', '', row_text)
        if "总计" in row_text_compact or "合计" in row_text_compact:
            continue
        if any(kw in row_text for kw in SKIP_HOUR_KEYWORDS):
            continue
        if hour_col < len(row.cells):
            val = row.cells[hour_col].text.strip()
            try:
                total += int(val)
            except ValueError:
                pass
    return total if total > 0 else None


def _check_hours(doc, findings: list[dict]) -> list[dict]:
    """课时验算主入口：先试 TP 模式，再试 CS 模式，最后兜底。"""
    if _hour_already_found(findings):
        return []

    # 策略1: TP 模式（表0=封面, 表1=进度表）
    result = _check_hours_tp(doc)
    if result:
        return result

    # 策略2: CS/EM 模式（表0=内容分配表，末行=总计）
    result = _check_hours_cs(doc)
    if result:
        return result

    # 如果CS模式成功验证（表0有学时总计行但无错误），说明学时匹配，不兜底
    if doc.tables:
        t0 = doc.tables[0]
        for row in t0.rows:
            row_text = " ".join(cell.text for cell in row.cells)
            if "学时总计" in row_text or "合 计" in row_text:
                return []

    # 策略3: 段落提取声明学时 + 任意表格累加（仅用于未知模板兜底）
    declared = _extract_hours_from_paragraphs(doc)
    if declared is not None:
        for t_idx, table in enumerate(doc.tables):
            table_sum = _sum_any_table(table)
            if table_sum is not None and table_sum >= declared * 0.3:
                if abs(table_sum - declared) > 2:
                    return [{
                        "rule": "hour_calculation_error",
                        "severity": "error",
                        "description": f"表格{t_idx + 1}各行学时累加得{table_sum}学时，但文档声明总学时为{declared}，相差{abs(table_sum - declared)}学时",
                        "location": f"表格{t_idx + 1}",
                        "suggestion": f"核实表格各行的学时数字，使累加总和等于{declared}",
                    }]

    return []


def _check_teacher_names(doc, findings: list[dict]) -> list[dict]:
    """教师姓名一致性检查。"""
    if _teacher_already_found(findings):
        return []

    occurrences: list[tuple[str, str, str]] = []

    for p_idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        for pattern in TEACHER_LABEL_PATTERNS:
            m = re.search(pattern, text)
            if m:
                name = m.group(1)
                if re.match(r'^[一-鿿]{2,4}$', name) and not any(c in name for c in _NON_NAME_CHARS):
                    occurrences.append((name, f"段落{p_idx + 1}", text[:60]))

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                text = cell.text.strip()
                for pattern in TEACHER_LABEL_PATTERNS:
                    m = re.search(pattern, text)
                    if m:
                        name = m.group(1)
                        if re.match(r'^[一-鿿]{2,4}$', name) and not any(c in name for c in _NON_NAME_CHARS):
                            occurrences.append((name, f"表格{t_idx + 1}第{r_idx + 1}行第{c_idx + 1}列", text[:60]))

    if len(occurrences) < 2:
        return []

    names = [n for n, _, _ in occurrences]
    unique_names = set(names)
    if len(unique_names) > 1:
        details = [f"{loc}出现「{name}」" for name, loc, _ in occurrences]
        return [{
            "rule": "teacher_mismatch",
            "severity": "error",
            "description": "教师姓名不一致：" + "；".join(details),
            "location": "、".join(loc for _, loc, _ in occurrences),
            "suggestion": "统一教师姓名写法",
        }]
    return []


def _check_school_name(doc, findings: list[dict]) -> list[dict]:
    """学校名检查：搜索旧校名。"""
    if _already_has_rule(findings, "school_name_error"):
        return []

    for p_idx, para in enumerate(doc.paragraphs):
        if WRONG_SCHOOL_NAME in para.text:
            return [{
                "rule": "school_name_error",
                "severity": "error",
                "description": f"学校名称错误：出现「{WRONG_SCHOOL_NAME}」，正确名称应为「{CORRECT_SCHOOL_NAME}」",
                "location": f"段落{p_idx + 1}",
                "suggestion": f"将「{WRONG_SCHOOL_NAME}」改为「{CORRECT_SCHOOL_NAME}」",
            }]

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if WRONG_SCHOOL_NAME in cell.text:
                    return [{
                        "rule": "school_name_error",
                        "severity": "error",
                        "description": f"学校名称错误：出现「{WRONG_SCHOOL_NAME}」，正确名称应为「{CORRECT_SCHOOL_NAME}」",
                        "location": f"表格{t_idx + 1}第{r_idx + 1}行第{c_idx + 1}列",
                        "suggestion": f"将「{WRONG_SCHOOL_NAME}」改为「{CORRECT_SCHOOL_NAME}」",
                    }]
    return []


def _check_known_typos(doc, findings: list[dict]) -> list[dict]:
    """已知错别字模式匹配。"""
    existing_descs = {f.get("description", "") for f in findings if f["rule"] == "typo"}

    new_findings: list[dict] = []

    for p_idx, para in enumerate(doc.paragraphs):
        text = para.text
        for wrong, correct in KNOWN_TYPOS:
            if wrong in text:
                if wrong not in str(existing_descs):
                    new_findings.append({
                        "rule": "typo",
                        "severity": "warning",
                        "description": f"发现错别字：「{wrong}」应为「{correct}」",
                        "location": f"段落{p_idx + 1}",
                        "suggestion": f"将「{wrong}」改为「{correct}」",
                    })

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                text = cell.text
                for wrong, correct in KNOWN_TYPOS:
                    if wrong in text:
                        if wrong not in str(existing_descs):
                            new_findings.append({
                                "rule": "typo",
                                "severity": "warning",
                                "description": f"发现错别字：「{wrong}」应为「{correct}」",
                                "location": f"表格{t_idx + 1}第{r_idx + 1}行第{c_idx + 1}列",
                                "suggestion": f"将「{wrong}」改为「{correct}」",
                            })
    return new_findings


def post_process_findings(filepath: str, ai_findings: list[dict], _extracted_info: dict) -> list[dict]:
    """AI 审查后的确定性后处理补丁。重新打开 docx 做结构化校验。"""
    try:
        doc = Document(filepath)
    except Exception:
        return ai_findings

    enriched = list(ai_findings)
    enriched.extend(_check_hours(doc, enriched))
    enriched.extend(_check_teacher_names(doc, enriched))
    enriched.extend(_check_school_name(doc, enriched))
    enriched.extend(_check_known_typos(doc, enriched))
    return enriched



def extract_doc_text(filepath: str) -> str:
    doc = Document(filepath)
    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for t_idx, table in enumerate(doc.tables):
        parts.append(f"\n[表格{t_idx + 1}]")
        for r_idx, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            parts.append(f"  行{r_idx + 1}: " + " | ".join(cells))

    return "\n".join(parts)


def build_prompt(doc_text: str, filename: str) -> str:
    text = doc_text[:15000] if len(doc_text) > 15000 else doc_text

    return f"""你是一个高校教务文件审查助手。请对以下教学文档进行全面审查，找出所有问题。

## 第一步：提取文档基本信息
从文档内容中提取：
- teacher: 教师姓名（通常出现在"任课教师""教师""授课教师"等字段旁，2-4字中文姓名）
- course: 课程名称（通常出现在"课程名称""课程""授课课程"等字段旁）
- class_name: 班级名称（如果有）

## 第二步：全面审查以下规则

### 规则1：学校名称错误（severity: error）
学校正确名称为"深圳信息职业技术大学"。检查是否出现：
- "深圳信息职业技术学院"（混淆旧校名，这是最常见的错误）
- 其他学校名称的错别字、遗漏、多余字

### 规则2：教师信息不匹配（severity: error）
检查文档中所有出现教师姓名的地方是否一致。教师姓名可能出现的位置包括：
- "任课教师"字段
- "教研室主任签名"字段
- "教师姓名"字段
- "编写执笔人"字段
- "审定负责人"字段
如果同一个教师的名字在不同位置写法不一致（如"曹维"vs"曹唯"、"张老师"vs"张教师"），标记为 teacher_mismatch。
注意：教师名字的错误写法不要归类为 typo，要归类为 teacher_mismatch。

### 规则3：错别字（severity: warning）
通过上下文识别明显错别字，如：
- "授科计划"应为"授课计划"
- "客程标准"应为"课程标准"
- 其他明显的中文错别字

### 规则4：课时计算校验（severity: error）—— 必须逐行数学验算！

你需要像一个会计查账一样，逐行核对课时数字。严格按以下步骤执行，不要跳过！

**步骤1 — 找到文档声明的总学时**
在文档开头段落中搜索"学时"或"总学时"或"计划教学时数"，记下这个数字。
例如："学时：64"、"总学时72"、"计划教学时数 64"

**步骤2 — 找到教学进度/内容表格**
找到包含具体教学内容的表格（特征：有"教学内容""学时""讲授""实践"等列标题的那个表格，不是封面信息表）。
如果文档有多个表格，教学进度表通常是包含多行教学单元的那个。

**步骤3 — 逐行提取学时数字并累加（关键！不要假设每行相同！）**
把教学进度表中每一行的"学时"列数字逐一读取出来，注意：每一行的学时数字可能不同，不要假设它们都一样！
把每一个数字逐个相加，得到实际累加总和。
注意：
- "学时"列通常在表格的第3列（序号之后）
- 排除标记为"复习""考试""机动""假期""节假日""国庆""五一"的行，但正常的教学周（即使内容包含"总结""复习"字样）应计入
- 每行只取一个学时数字

**步骤4 — 交叉验证**
同时检查以下所有条件：
- 进度表累加总和 是否等于 步骤1找到的总学时？不等 → hour_calculation_error
- 如果有汇总表写着"周数×周学时"：乘积 是否等于 总学时？不等 → hour_calculation_error
- 如果同一个数字在两个地方出现但值不同 → hour_inconsistency

**步骤5 — 输出计算过程**
在 description 中必须写清楚你找到的具体数字和计算过程。
示例格式："文档声明总学时64，教学进度表16行逐行累加得68（4+8+4+4+4+4+4+4+4+4+4+4+4+4+4+4=68），比声明多4学时"

**重要：只在数字明显对不上时报错**
- 如果累加结果与声明值相差在2学时以内，可能是统计口径差异，不报错
- 如果排除了非教学行后确实对不上，才标记为错误

### 规则5：格式问题（severity: warning）
检查文档格式是否规范：

5a. 字体规范
- 正文应使用宋体、仿宋、楷体或黑体等常见中文字体
- 如果出现不常见的字体（如艺术字体、英文字体用于中文正文），标记为格式问题
- 同一文档中使用的字体种类不应超过4种

5b. 字号规范
- 正文标题字号是否过大或过小（正文建议小四/12pt，标题建议三号/16pt到小二/18pt）
- 同一层级标题字号是否一致

5c. 对齐方式
- 正文是否两端对齐或左对齐
- 标题是否居中
- 表格内容是否对齐一致

## 重要提醒 —— 输出前必须逐条核对！
在生成JSON之前，请在心里完成以下检查清单，确保没有遗漏：

[ ] 学校名：我扫描了全文，确认了"深圳信息职业技术大学"是否被写错（如写成"职业技术学院"）？有无结果都要报告。
[ ] 教师名：我找到了文档中所有出现教师姓名的地方（任课教师、教研室主任、编写执笔人、审定负责人等），对比是否一致？例如p[10]"曹维"和p[12]"曹维"必须一致，不一致=teacher_mismatch。教师名字的错误写法（如"曹为""曹唯""陈叶羽各"）不要归类为typo，要归类为teacher_mismatch。
[ ] 错别字：我逐段扫描了是否有明显中文错别字（授科计划、客程标准、职叶能力、核心技木、数锯分析、产叶、人工知能、课成、能立、基楚、人材培养、线形代数、距阵、统记学等）？
[ ] 课时：我逐行提取了教学进度表中每一行的学时数字并累加，与声明总学时对比？重要：每一行的学时可能不同，必须逐个读取，不要假设所有行都是相同数字！
[ ] 格式：我检查了字体、字号、对齐是否规范？

记住：发现课时错误后，继续检查教师名和错别字，不能跳过！每条规则独立检查。

## 输出格式
在输出之前，请确认你已经检查了所有5条规则。如果某条规则没有发现问题，就不要在findings中添加该规则的条目。
严格按以下JSON格式输出，不要输出任何其他内容：
```json
{{
  "extracted_info": {{
    "teacher": "教师姓名或null",
    "course": "课程名称或null",
    "class_name": "班级名称或null"
  }},
  "findings": [
    {{
      "rule": "school_name_error",
      "severity": "error",
      "description": "详细描述问题，包含原文和正确写法",
      "location": "表格编号/段落/行号",
      "suggestion": "建议修改方案"
    }},
    {{
      "rule": "teacher_mismatch",
      "severity": "error",
      "description": "详细描述不匹配的具体情况",
      "location": "问题位置",
      "suggestion": "建议"
    }},
    {{
      "rule": "typo",
      "severity": "warning",
      "description": "错别字及正确写法",
      "location": "错别字位置",
      "suggestion": "建议修改为..."
    }},
    {{
      "rule": "hour_calculation_error",
      "severity": "error",
      "description": "详细描述课时计算问题，包含你做的具体数学计算过程，如累加公式",
      "location": "表格编号",
      "suggestion": "建议核实并修正"
    }},
    {{
      "rule": "hour_inconsistency",
      "severity": "error",
      "description": "进度表累加与汇总表不一致的具体数字",
      "location": "表格编号",
      "suggestion": "建议核实并统一"
    }},
    {{
      "rule": "format_issue",
      "severity": "warning",
      "description": "详细描述格式问题，包含当前状态和规范要求",
      "location": "段落/表格位置",
      "suggestion": "建议调整为..."
    }}
  ]
}}
```

如果没有发现任何问题，返回 {{"extracted_info": {{"teacher": null, "course": null, "class_name": null}}, "findings": []}}。

## 文档文件名
{filename}

## 文档内容
{text}"""


def parse_ai_response(raw_text: str) -> tuple[list[dict], dict]:
    json_match = re.search(r"```(?:json)?\s*([\s\S]*?)```", raw_text)
    if json_match:
        raw_text = json_match.group(1).strip()

    data = json.loads(raw_text)
    findings = data.get("findings", [])
    extracted_info = data.get("extracted_info", {}) or {}

    valid_rules = {"school_name_error", "teacher_mismatch", "typo", "hour_calculation_error", "hour_inconsistency", "format_issue"}
    for f in findings:
        if f.get("rule") not in valid_rules:
            f["rule"] = "typo"
        f.setdefault("description", "无描述")
        f.setdefault("location", "")
        f.setdefault("suggestion", "")
        if "severity" not in f:
            if f["rule"] == "typo":
                f["severity"] = "warning"
            else:
                f["severity"] = "error"

    return findings, extracted_info


_proxy_handler = urllib.request.ProxyHandler({})
_opener = urllib.request.build_opener(_proxy_handler)


def _call_api(api_key: str, prompt: str) -> str:
    body = json.dumps({
        "model": MODEL_ID,
        "messages": [
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.1,
    }).encode("utf-8")

    req = urllib.request.Request(
        ARK_URL,
        data=body,
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
    )

    try:
        with _opener.open(req, timeout=300) as resp:
            result = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"AI API 返回错误 {e.code}: {e.read().decode('utf-8', errors='replace')[:500]}")
    except urllib.error.URLError as e:
        raise RuntimeError(f"无法连接 AI 服务: {e.reason}")

    return result["choices"][0]["message"]["content"]


async def review(filepath: str, filename: str, api_key: str | None = None) -> tuple[list[dict], dict, str]:
    key = api_key or ARK_KEY
    doc_text = extract_doc_text(filepath)
    prompt = build_prompt(doc_text, filename)
    content = await asyncio.to_thread(_call_api, key, prompt)
    findings, extracted_info = parse_ai_response(content)
    findings = post_process_findings(filepath, findings, extracted_info)
    return findings, extracted_info, content


async def _call_api_stream(api_key: str, prompt: str):
    body = {
        "model": MODEL_ID,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.1,
        "stream": True,
    }

    async with httpx.AsyncClient(timeout=httpx.Timeout(300.0)) as client:
        async with client.stream(
            "POST", ARK_URL,
            json=body,
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
            },
        ) as response:
            if response.status_code != 200:
                body_text = await response.aread()
                raise RuntimeError(f"AI API 返回 {response.status_code}: {body_text.decode('utf-8', errors='replace')[:500]}")

            async for line in response.aiter_lines():
                if not line.startswith("data: "):
                    continue
                data_str = line[6:]
                if data_str == "[DONE]":
                    return
                try:
                    chunk = json.loads(data_str)
                    delta = chunk.get("choices", [{}])[0].get("delta", {})
                    reasoning = delta.get("reasoning_content", "")
                    content = delta.get("content", "")
                    if reasoning:
                        yield ("reasoning", reasoning)
                    if content:
                        yield ("token", content)
                except json.JSONDecodeError:
                    continue


async def review_stream(filepath: str, filename: str, api_key: str | None = None):
    key = api_key or ARK_KEY
    doc_text = extract_doc_text(filepath)
    prompt = build_prompt(doc_text, filename)

    full_text_parts: list[str] = []
    try:
        async for event_type, token in _call_api_stream(key, prompt):
            if event_type == "reasoning":
                yield ("reasoning", token)
            elif event_type == "token":
                full_text_parts.append(token)
                yield ("token", token)
    except Exception as e:
        yield ("error", str(e))
        return

    full_text = "".join(full_text_parts)
    try:
        findings, extracted_info = parse_ai_response(full_text)
        findings = post_process_findings(filepath, findings, extracted_info)
        yield ("done", (findings, extracted_info, full_text))
    except Exception as e:
        yield ("error", f"解析 AI 响应失败: {e}")
