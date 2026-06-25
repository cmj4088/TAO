"""核心审查逻辑：格式检查
改编自 docs/document_content_extraction/review_by_template.py
"""
import os
import re
from collections import defaultdict

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.schemas.app02 import ReviewIssue, FileReviewResult

BASE = os.path.join(os.path.dirname(__file__), "..", "..", "..", "docs", "app02.file")
TEMPLATE_DIR = os.path.join(
    BASE, "附件4：教学文件（课程标准、计划、教案）",
    "附件4：教学文件（课程标准、计划、教案）模板",
)

TEMPLATES = {
    "课程标准": os.path.join(TEMPLATE_DIR, "附件1：课程标准（体例）.docx"),
    "授课计划": os.path.join(TEMPLATE_DIR, "附件2：学期授课计划表（体例）.docx"),
    "教案": os.path.join(TEMPLATE_DIR, "附件3：教案（体例）.docx"),
}

_template_cache: dict[str, dict] = {}


def classify_template(filename: str) -> str:
    if "课程标准" in filename:
        return "课程标准"
    if "授课计划" in filename:
        return "授课计划"
    if "教案" in filename:
        return "教案"
    return "未知"


def extract_teacher_from_filename(filename: str) -> str:
    name = os.path.splitext(filename)[0]
    parts = re.split(r'[_\-\s（）()《》]', name)
    for part in reversed(parts):
        if re.match(r'^[一-龥]{2,4}$', part) and part not in ("授课计划", "课程标准", "教案", "体例"):
            return part
    return "未知"


def _get_alignment_name(p) -> str:
    m = {
        WD_ALIGN_PARAGRAPH.LEFT: "左对齐",
        WD_ALIGN_PARAGRAPH.CENTER: "居中",
        WD_ALIGN_PARAGRAPH.RIGHT: "右对齐",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "两端对齐",
    }
    return m.get(p.alignment, "默认")


def extract_template_profile(template_path: str) -> dict:
    if template_path in _template_cache:
        return _template_cache[template_path]

    if not os.path.exists(template_path):
        return {"允许的字体": set(), "允许的字号": set(), "允许的对齐": set()}

    doc = Document(template_path)
    allowed_fonts: set[str] = set()
    allowed_sizes: set[str] = set()
    allowed_alignments: set[str] = set()

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    allowed_alignments.add(_get_alignment_name(para))
                    for run in para.runs:
                        if run.font.name:
                            allowed_fonts.add(run.font.name)
                        if run.font.size:
                            allowed_sizes.add(f"{run.font.size.pt:.0f}pt")

    profile = {
        "允许的字体": allowed_fonts,
        "允许的字号": allowed_sizes,
        "允许的对齐": allowed_alignments,
    }
    _template_cache[template_path] = profile
    return profile


def review_format(filepath: str, template_profile: dict) -> tuple[list[ReviewIssue], set[str]]:
    doc = Document(filepath)
    allowed_fonts = template_profile.get("允许的字体", set())
    allowed_sizes = template_profile.get("允许的字号", set())
    allowed_alignments = template_profile.get("允许的对齐", set())

    issues: list[ReviewIssue] = []
    all_fonts_used: set[str] = set()
    font_locations: dict[str, list[dict]] = defaultdict(list)
    alignment_locations: dict[str, list[dict]] = defaultdict(list)

    for t_idx, table in enumerate(doc.tables):
        t_key = f"表格{t_idx + 1}"

        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                text = cell.text.strip()
                if not text:
                    continue

                for para in cell.paragraphs:
                    align_name = _get_alignment_name(para)
                    if allowed_alignments and align_name not in allowed_alignments and align_name != "默认":
                        alignment_locations[align_name].append({
                            "表格": t_key,
                            "位置": f"第{r_idx + 1}行第{c_idx + 1}列",
                            "内容": text[:30],
                        })

                    for run in para.runs:
                        if not run.text.strip():
                            continue

                        fn = run.font.name or "默认字体"
                        fs = f"{run.font.size.pt:.0f}pt" if run.font.size else "默认字号"
                        all_fonts_used.add(fn)

                        if fn not in allowed_fonts and fn != "默认字体":
                            font_locations[fn].append({
                                "表格": t_key,
                                "位置": f"第{r_idx + 1}行第{c_idx + 1}列",
                                "内容": text[:30],
                            })

                        if fs != "默认字号" and fs not in allowed_sizes:
                            issues.append(ReviewIssue(
                                type="字号偏离",
                                detail=f"字号 {fs} 不在模板允许范围 {sorted(allowed_sizes)}",
                                location=f"{t_key} 第{r_idx + 1}行第{c_idx + 1}列",
                                snippet=text[:30],
                                severity="warning",
                            ))

    if font_locations:
        for font, locs in font_locations.items():
            unique_locs: list[str] = []
            seen: set[tuple] = set()
            for loc in locs:
                key = (loc["表格"], loc["位置"])
                if key not in seen:
                    unique_locs.append(f"{loc['表格']}{loc['位置']}")
                    seen.add(key)
                    if len(unique_locs) >= 3:
                        break
            examples = "、".join(unique_locs)
            issues.append(ReviewIssue(
                type="字体偏离",
                detail=f"使用了模板不允许的字体「{font}」，共{len(locs)}处",
                location=examples,
                snippet="",
                severity="warning",
            ))

    if alignment_locations:
        for align, locs in alignment_locations.items():
            unique_locs: list[str] = []
            seen: set[tuple] = set()
            for loc in locs:
                key = (loc["表格"], loc["位置"])
                if key not in seen:
                    unique_locs.append(f"{loc['表格']}{loc['位置']}")
                    seen.add(key)
                    if len(unique_locs) >= 3:
                        break
            examples = "、".join(unique_locs)
            issues.append(ReviewIssue(
                type="对齐偏离",
                detail=f"使用了模板不允许的对齐方式「{align}」，共{len(locs)}处",
                location=examples,
                snippet="",
                severity="warning",
            ))

    if len(all_fonts_used) > 5:
        issues.append(ReviewIssue(
            type="字体过多",
            detail=f"全文使用了{len(all_fonts_used)}种字体（上限5种）: {', '.join(sorted(all_fonts_used))}",
            location="全文",
            snippet="",
            severity="warning",
        ))

    return issues, all_fonts_used


def review_single(filepath: str, filename: str) -> FileReviewResult:
    file_type = classify_template(filename)
    teacher = extract_teacher_from_filename(filename)

    if file_type not in TEMPLATES:
        return FileReviewResult(
            file_id="",
            filename=filename,
            file_type=file_type,
            teacher=teacher,
            passed=True,
            fonts_used=[],
            issues=[],
        )

    template_path = TEMPLATES[file_type]
    if not os.path.exists(template_path):
        return FileReviewResult(
            file_id="",
            filename=filename,
            file_type=file_type,
            teacher=teacher,
            passed=True,
            fonts_used=[],
            issues=[],
        )

    profile = extract_template_profile(template_path)
    format_issues, fonts_used = review_format(filepath, profile)

    return FileReviewResult(
        file_id="",
        filename=filename,
        file_type=file_type,
        teacher=teacher,
        passed=True,
        fonts_used=sorted(fonts_used),
        issues=format_issues,
    )
